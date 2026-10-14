import io
import re
import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import fitz
import qrcode

st.set_page_config(page_title="Advanced Document Formatter", page_icon="📄")
st.title("📄 Advanced PDF / Word Formatter")
st.caption("Page borders, per-page headers/footers, page numbers, preview, payment, and download.")

# ---------------- Payment settings ----------------
UPI_ID = "7676331443-2@ybl"       # Change this
PAYEE_NAME = "Rohith"      # Change this
AMOUNT = 10                   # Amount in INR

uploaded = st.file_uploader("Upload a Word or PDF file", type=["docx", "pdf"])

# ---------------- Sidebar: text formatting ----------------
st.sidebar.header("Text formatting")

font_name = st.sidebar.text_input("Font name", "Arial")
font_size = st.sidebar.number_input("Font size", 8, 40, 12)
bold = st.sidebar.checkbox("Bold")
italic = st.sidebar.checkbox("Italic")
underline = st.sidebar.checkbox("Underline")

alignment = st.sidebar.selectbox(
    "Body text alignment",
    ["Left", "Center", "Right", "Justify"]
)

text_color = st.sidebar.color_picker("Body text color", "#000000")

# ---------------- Sidebar: page border ----------------
st.sidebar.subheader("Page border")

add_border = st.sidebar.checkbox("Add page border", False)

border_style = st.sidebar.selectbox(
    "Border design",
    [
        "Solid",
        "Dashed",
        "Dotted",
        "Double",
        "Thick Outline",
        "Corner Accent"
    ]
)

border_color = st.sidebar.color_picker("Border color", "#000000")
border_width = st.sidebar.slider("Border width (PDF points)", 0.5, 8.0, 1.5, 0.5)
border_margin = st.sidebar.slider("Border margin (PDF points)", 10, 60, 25, 5)
# ---------------- Sidebar: header/footer ----------------
st.sidebar.subheader("Header / Footer")

header_text = st.sidebar.text_input("Default header text", "")
footer_text = st.sidebar.text_input("Default footer text", "")

header_align = st.sidebar.selectbox(
    "Header alignment",
    ["Left", "Center", "Right"]
)

footer_align = st.sidebar.selectbox(
    "Footer alignment",
    ["Left", "Center", "Right"]
)

header_color = st.sidebar.color_picker("Header color", "#808080")
footer_color = st.sidebar.color_picker("Footer color", "#808080")

header_size = st.sidebar.number_input("Header font size", 6, 20, 8)
footer_size = st.sidebar.number_input("Footer font size", 6, 20, 8)

st.sidebar.markdown("**Custom per-page text**")
st.sidebar.caption(
    "Enter one line per page. Leave blank to use the default header/footer.\n\n"
    "Example for 3 pages:\n"
    "`Page 1 header`\n"
    "`Page 2 header`\n"
    "`Page 3 header`"
)

custom_headers = st.sidebar.text_area("Custom headers", height=100)
custom_footers = st.sidebar.text_area("Custom footers", height=100)

# ---------------- Sidebar: page numbers ----------------
st.sidebar.subheader("Page numbers")

add_page_numbers = st.sidebar.checkbox("Add page numbers", True)

page_number_position = st.sidebar.selectbox(
    "Page number position",
    ["Left", "Center", "Right"]
)

page_number_format = st.sidebar.selectbox(
    "Page number format",
    ["Page X of Y", "X of Y", "X", "- X -"]
)

start_page_number = st.sidebar.number_input(
    "Start page number from",
    1,
    10000,
    1
)

# ---------------- Maps ----------------
ALIGN_MAP = {
    "Left": WD_ALIGN_PARAGRAPH.LEFT,
    "Center": WD_ALIGN_PARAGRAPH.CENTER,
    "Right": WD_ALIGN_PARAGRAPH.RIGHT,
    "Justify": WD_ALIGN_PARAGRAPH.JUSTIFY
}

PDF_ALIGN_MAP = {
    "Left": fitz.TEXT_ALIGN_LEFT,
    "Center": fitz.TEXT_ALIGN_CENTER,
    "Right": fitz.TEXT_ALIGN_RIGHT,
    "Justify": fitz.TEXT_ALIGN_JUSTIFY
}


def hex_to_rgb01(hex_color):
    return tuple(int(hex_color[i:i + 2], 16) / 255 for i in (1, 3, 5))


def get_custom_lines(text):
    return [line.strip() for line in text.splitlines()] if text.strip() else []


def get_page_text(custom_lines, default_text, page_index):
    if page_index < len(custom_lines) and custom_lines[page_index]:
        return custom_lines[page_index]
    return default_text


def make_page_number(current, total):
    display_number = current + start_page_number - 1

    if page_number_format == "Page X of Y":
        return f"Page {display_number} of {total}"
    elif page_number_format == "X of Y":
        return f"{display_number} of {total}"
    elif page_number_format == "X":
        return str(display_number)
    else:
        return f"- {display_number} -"


# ---------------- DOCX helpers ----------------
def add_docx_page_border(section):
    sect_pr = section._sectPr

    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")

    r, g, b = tuple(int(border_color[i:i + 2], 16) for i in (1, 3, 5))
    color_hex = f"{r:02X}{g:02X}{b:02X}"

    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(int(border_width * 8)))
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), color_hex)
        pg_borders.append(border)

    sect_pr.append(pg_borders)


def add_docx_page_number(paragraph):
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instruction)
    run._r.append(fld_end)


# ---------------- DOCX formatter ----------------
def format_docx(file):
    doc = Document(file)

    for paragraph in doc.paragraphs:
        paragraph.alignment = ALIGN_MAP[alignment]

        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.bold = bold
            run.italic = italic
            run.underline = underline

            r, g, b = tuple(int(text_color[i:i + 2], 16) for i in (1, 3, 5))
            run.font.color.rgb = RGBColor(r, g, b)

    for section in doc.sections:
        if add_border:
            add_docx_page_border(section)

        header_paragraph = section.header.paragraphs[0]
        footer_paragraph = section.footer.paragraphs[0]

        header_paragraph.text = header_text
        footer_paragraph.text = footer_text

        header_paragraph.alignment = ALIGN_MAP[header_align]
        footer_paragraph.alignment = ALIGN_MAP[footer_align]

        for run in header_paragraph.runs:
            run.font.size = Pt(header_size)
            run.font.name = font_name

            r, g, b = tuple(int(header_color[i:i + 2], 16) for i in (1, 3, 5))
            run.font.color.rgb = RGBColor(r, g, b)

        for run in footer_paragraph.runs:
            run.font.size = Pt(footer_size)
            run.font.name = font_name

            r, g, b = tuple(int(footer_color[i:i + 2], 16) for i in (1, 3, 5))
            run.font.color.rgb = RGBColor(r, g, b)

        if add_page_numbers:
            footer_paragraph.add_run("\t\t")
            add_docx_page_number(footer_paragraph)
            footer_paragraph.alignment = ALIGN_MAP[footer_align]

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


# ---------------- PDF formatter ----------------
def format_pdf(file):
    doc = fitz.open(stream=file.read(), filetype="pdf")

    total_pages = doc.page_count
    custom_header_lines = get_custom_lines(custom_headers)
    custom_footer_lines = get_custom_lines(custom_footers)

    header_rgb = hex_to_rgb01(header_color)
    footer_rgb = hex_to_rgb01(footer_color)
    border_rgb = hex_to_rgb01(border_color)

    for page in doc:
        width = page.rect.width
        height = page.rect.height
        page_index = page.number

        # Page border
        if add_border:
            page.draw_rect(
                fitz.Rect(
                    border_margin,
                    border_margin,
                    width - border_margin,
                    height - border_margin
                ),
                color=border_rgb,
                width=border_width
            )

        # Header
        current_header = get_page_text(
            custom_header_lines,
            header_text,
            page_index
        )

        if current_header:
            page.insert_textbox(
                fitz.Rect(30, 8, width - 30, 30),
                current_header,
                fontname="helv",
                fontsize=header_size,
                color=header_rgb,
                align=PDF_ALIGN_MAP[header_align]
            )

        # Footer text
        current_footer = get_page_text(
            custom_footer_lines,
            footer_text,
            page_index
        )

        # Page number
        page_number_text = ""
        if add_page_numbers:
            page_number_text = make_page_number(page_index, total_pages)

        # Combine footer and page number
        if current_footer and page_number_text:
            footer_content = f"{current_footer}   |   {page_number_text}"
        elif current_footer:
            footer_content = current_footer
        else:
            footer_content = page_number_text

        if footer_content:
            page.insert_textbox(
                fitz.Rect(30, height - 26, width - 30, height - 8),
                footer_content,
                fontname="helv",
                fontsize=footer_size,
                color=footer_rgb,
                align=PDF_ALIGN_MAP[footer_align]
            )

    result = doc.tobytes()
    doc.close()
    return result


# ---------------- UPI QR ----------------
def make_upi_qr(amount):
    upi_link = (
        f"upi://pay?pa={UPI_ID}"
        f"&pn={PAYEE_NAME.replace(' ', '%20')}"
        f"&am={amount}"
        f"&cu=INR"
        f"&tn=Document%20Formatting"
    )

    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(upi_link)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")

    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer.getvalue()


# ---------------- Main app ----------------
if uploaded:
    st.subheader("⚙️ Format document")

    if st.button("Format document"):
        try:
            uploaded.seek(0)

            if uploaded.name.lower().endswith(".docx"):
                formatted = format_docx(uploaded)
                st.session_state["formatted_file"] = formatted
                st.session_state["download_name"] = "formatted_document.docx"
                st.session_state["mime"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                st.session_state["file_type"] = "docx"

            else:
                formatted = format_pdf(uploaded)
                st.session_state["formatted_file"] = formatted
                st.session_state["download_name"] = "formatted_document.pdf"
                st.session_state["mime"] = "application/pdf"
                st.session_state["file_type"] = "pdf"

            st.success("Document formatted successfully!")

        except Exception as e:
            st.error(f"Could not format this file: {e}")

    # ---------------- Preview after formatting ----------------
    if "formatted_file" in st.session_state:
        st.subheader("👁️ Formatted preview")

        if st.session_state["file_type"] == "pdf":
            formatted_doc = fitz.open(
                stream=st.session_state["formatted_file"],
                filetype="pdf"
            )

            preview_page = st.number_input(
                "Preview page",
                1,
                formatted_doc.page_count,
                1
            ) - 1

            pix = formatted_doc[preview_page].get_pixmap(dpi=120)

            st.image(
                pix.tobytes("png"),
                caption=f"Formatted page {preview_page + 1}",
                use_container_width=True
            )

            formatted_doc.close()

        else:
            try:
                doc = Document(io.BytesIO(st.session_state["formatted_file"]))

                preview_text = "\n".join(
                    p.text for p in doc.paragraphs if p.text.strip()
                )

                st.text_area(
                    "Formatted Word text preview",
                    preview_text[:3000],
                    height=300
                )

                st.info(
                    "DOCX preview shows text content. "
                    "Download the file to view the border, header, footer, and page number in Word."
                )

            except Exception as e:
                st.error(f"Could not preview the formatted Word file: {e}")

        # ---------------- Payment ----------------
        st.subheader("💳 Payment")

        col1, col2 = st.columns([1, 1])

        with col1:
            st.markdown(f"**Amount to pay:** ₹{AMOUNT}")
            st.markdown(f"**UPI ID:** `{UPI_ID}`")
            st.markdown("Scan the QR using GPay, PhonePe, Paytm, or any UPI app.")

        with col2:
            qr_image = make_upi_qr(AMOUNT)
            st.image(qr_image, caption="Scan to pay", width=220)

        st.divider()

        paid = st.checkbox("I have completed the payment")

        # ---------------- Download ----------------
        if paid:
            st.subheader("⬇️ Download")

            st.download_button(
                label="Download formatted file",
                data=st.session_state["formatted_file"],
                file_name=st.session_state["download_name"],
                mime=st.session_state["mime"]
            )

        else:
            st.warning("Complete the payment and tick the confirmation box to unlock download.")

else:
    st.info("Upload a `.docx` or `.pdf` file to begin.")